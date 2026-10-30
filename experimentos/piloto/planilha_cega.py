"""
Planilha de anotação CEGA do piloto: substitui a planilha de preanotar.py na anotação humana.

O que fica escondido de quem anota:
- a pré-anotação heurística e o veredito do verificador (a anotação humana é a referência, D029/D034);
- o modelo e a repetição (as respostas dos arquivos são misturadas em ordem aleatória, com semente fixa).
Cada resposta recebe um código opaco (A001, A002…). A correspondência código → arquivo/modelo/vinheta/
repetição vai para um arquivo-chave separado, que NÃO deve ser aberto antes do fim da anotação.

Limite (declarar): o estilo de escrita pode denunciar o modelo; a cegueira é quanto ao rótulo, não ao estilo.

Uso:
  python3 planilha_cega.py gerar resultados/<a>.jsonl resultados/<b>.jsonl [--repeticao 1]
      → resultados/anotacao_cega_<carimbo>.xlsx  e  resultados/anotacao_cega_<carimbo>_CHAVE.csv
  python3 planilha_cega.py converter resultados/anotacao_cega_<carimbo>.csv   (planilha antiga em CSV → .xlsx)
  python3 planilha_cega.py revelar resultados/anotacao_cega_<carimbo>.xlsx
      → resultados/anotacao_cega_<carimbo>_revelada.csv  (formato aceito por resumir.py)
"""

import argparse
import csv
import datetime
import json
import random
from pathlib import Path

SEMENTE = 20261026
COLUNAS_H = ["H_totalmente_fiel(S/N)", "H_erros_polaridade", "H_erros_numericos",
             "H_fora_do_modelo", "H_nao_sustentadas", "H_observacoes"]
FORA_DO_MODELO = ("imc", "raca_cor")  # entradas do cadastro que o PREVENT-ASCVD não usa
FRACAO_EPSILON = 0.05  # ε = 5% de Σ|φ|: contribuição desprezível (metricas-fidelidade.md §3.3)
NOMES = {
    "idade": "idade", "pas": "pressão sistólica", "nao_hdl": "colesterol não-HDL", "hdl": "HDL (colesterol bom)",
    "tfge": "função renal (TFGe)", "anti_hipertensivo": "uso de remédio para pressão",
    "estatina": "uso de estatina", "diabetes": "diabetes", "tabagismo": "tabagismo",
}


def _direcoes(atribuicao):
    """Gabarito de polaridade em palavras: só reorganiza a atribuição, não julga a resposta."""
    eps = FRACAO_EPSILON * sum(abs(v) for v in atribuicao.values())
    ordem = sorted(atribuicao.items(), key=lambda kv: -abs(kv[1]))
    grupo = lambda cond: ", ".join(NOMES[k] for k, v in ordem if cond(v)) or "nenhum"
    return ("AUMENTARAM: %s || DIMINUÍRAM: %s || DESPREZÍVEIS (qualquer direção vale): %s"
            % (grupo(lambda v: v >= eps), grupo(lambda v: v <= -eps), grupo(lambda v: abs(v) < eps)))


def _referencia(g):
    """Pacote de referência em texto, para conferir a resposta sem abrir o gabarito.json."""
    e = g["entradas"]
    entradas = ", ".join("%s=%s" % (k, v) for k, v in e.items() if k not in FORA_DO_MODELO)
    fora = ", ".join("%s=%s" % (k, e[k]) for k in FORA_DO_MODELO if k in e)
    atrib = sorted(g["atribuicao"].items(), key=lambda kv: -abs(kv[1]))
    return {
        "risco_pct": g["risco_pct"],
        "categoria": g["categoria"],
        "fatores(gabarito)": _direcoes(g["atribuicao"]),
        "regras_aplicadas": " | ".join(g.get("regras_aplicadas", [])),
        "metas": "; ".join("%s %s" % kv for kv in g["metas"].items()),
        "atribuicao(ordem |φ|)": "; ".join("%s %+.4f" % kv for kv in atrib),
        "entradas_do_modelo": entradas,
        "fora_do_modelo": fora,
    }


# Ordem na planilha: gabarito à esquerda, resposta no meio, veredito logo depois; detalhes no fim.
ORDEM = ["codigo", "vinheta", "risco_pct", "categoria", "fatores(gabarito)", "metas", "regras_aplicadas",
         "fora_do_modelo", "resposta"] + COLUNAS_H + ["atribuicao(ordem |φ|)", "entradas_do_modelo"]
LARGURAS = {"codigo": 7, "vinheta": 7, "risco_pct": 7, "categoria": 12, "fatores(gabarito)": 45, "metas": 22,
            "regras_aplicadas": 30, "fora_do_modelo": 18, "resposta": 80, "H_totalmente_fiel(S/N)": 10,
            "H_observacoes": 35, "atribuicao(ordem |φ|)": 40, "entradas_do_modelo": 40}
ALTURA_LINHA = 75  # pontos (~5 linhas de texto); o resto aparece ao clicar na célula


def salvar_xlsx(linhas, caminho):
    """Planilha de anotação em .xlsx: sem tela de importação (o CSV quebrava colunas no LibreOffice),
    altura de linha fixa e S/N validado. Dependência: openpyxl (só para este script)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "anotacao"
    ws.append(ORDEM)
    for linha in linhas:
        ws.append([linha.get(c, "") for c in ORDEM])
    for row in ws.iter_rows(min_row=2):
        for cel in row:
            if isinstance(cel.value, str) and cel.value[:1] in "=+-@":
                cel.data_type = "s"  # texto, nunca fórmula
            cel.alignment = Alignment(wrap_text=True, vertical="top")
    destaque = PatternFill("solid", fgColor="FFF2CC")
    for j, nome in enumerate(ORDEM, 1):
        letra = get_column_letter(j)
        ws.column_dimensions[letra].width = LARGURAS.get(nome, 14)
        cab = ws.cell(row=1, column=j)
        cab.font = Font(bold=True)
        cab.alignment = Alignment(wrap_text=True, vertical="top")
        if nome.startswith("H_"):
            cab.fill = destaque
            for i in range(2, len(linhas) + 2):
                ws.cell(row=i, column=j).fill = destaque
    for i in range(2, len(linhas) + 2):
        ws.row_dimensions[i].height = ALTURA_LINHA
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    col_sn = get_column_letter(ORDEM.index("H_totalmente_fiel(S/N)") + 1)
    dv = DataValidation(type="list", formula1='"S,N"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("%s2:%s%d" % (col_sn, col_sn, len(linhas) + 1))
    wb.save(caminho)


def ler_planilha(caminho):
    if caminho.suffix == ".xlsx":
        from openpyxl import load_workbook
        ws = load_workbook(caminho, read_only=True).active
        linhas = ws.iter_rows(values_only=True)
        cab = next(linhas)
        return [{k: ("" if v is None else str(v)) for k, v in zip(cab, l)} for l in linhas if any(l)]
    with caminho.open(encoding="utf-8-sig") as f:
        return list(csv.DictReader(f, delimiter=";"))


def gerar(arquivos, repeticao=None):
    dir_ = arquivos[0].parent
    gabs = {g["id"]: g for g in json.loads((dir_ / "gabarito.json").read_text(encoding="utf-8"))}
    registros, descartados = [], 0
    for arq in arquivos:
        for linha in arq.read_text(encoding="utf-8").splitlines():
            r = json.loads(linha)
            if repeticao is not None and r["repeticao"] != repeticao:
                continue
            if r.get("erro") or not r.get("resposta"):
                descartados += 1
                continue
            registros.append((arq.name, r))
    random.Random(SEMENTE).shuffle(registros)

    carimbo = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    planilha = dir_ / ("anotacao_cega_%s.xlsx" % carimbo)
    chave = dir_ / ("anotacao_cega_%s_CHAVE.csv" % carimbo)
    linhas, linhas_chave = [], []
    for i, (nome, r) in enumerate(registros, 1):
        codigo = "A%03d" % i
        linhas.append(dict({"codigo": codigo, "vinheta": r["vinheta"]}, **_referencia(gabs[r["vinheta"]]),
                           resposta=r["resposta"], **{c: "" for c in COLUNAS_H}))
        linhas_chave.append({"codigo": codigo, "arquivo": nome, "modelo": r["modelo"], "vinheta": r["vinheta"],
                             "repeticao": r["repeticao"], "commit": r.get("commit", "")})
    salvar_xlsx(linhas, planilha)
    with chave.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(linhas_chave[0].keys()), delimiter=";")
        w.writeheader()
        w.writerows(linhas_chave)
    print("%d respostas na planilha cega (%d descartadas por erro ou resposta vazia)" % (len(linhas), descartados))
    print("Anotar:      %s" % planilha)
    print("NÃO abrir:   %s" % chave)


def converter(csv_antigo):
    """Converte uma planilha cega já gerada em CSV para .xlsx, mantendo os códigos (a chave continua valendo)."""
    destino = csv_antigo.with_suffix(".xlsx")
    salvar_xlsx(ler_planilha(csv_antigo), destino)
    print(destino)


def revelar(planilha):
    chave = planilha.with_name(planilha.stem + "_CHAVE.csv")
    with chave.open(encoding="utf-8-sig") as f:
        mapa = {l["codigo"]: l for l in csv.DictReader(f, delimiter=";")}
    anotadas = ler_planilha(planilha)
    saida = planilha.with_name(planilha.stem + "_revelada.csv")
    campos = ["codigo", "modelo", "vinheta", "repeticao", "arquivo"] + COLUNAS_H + ["resposta"]
    with saida.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=campos, delimiter=";", extrasaction="ignore")
        w.writeheader()
        for l in anotadas:
            w.writerow(dict(l, **mapa[l["codigo"]]))
    pendentes = sum(1 for l in anotadas if l["H_totalmente_fiel(S/N)"].strip().upper() not in ("S", "N"))
    print("%s (%d linhas; %d sem H_totalmente_fiel preenchido)" % (saida, len(anotadas), pendentes))


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    g = sub.add_parser("gerar")
    g.add_argument("jsonl", nargs="+", type=Path)
    g.add_argument("--repeticao", type=int, help="só esta repetição (D035: o piloto anota a 1ª de cada vinheta)")
    c = sub.add_parser("converter")
    c.add_argument("csv", type=Path)
    r = sub.add_parser("revelar")
    r.add_argument("planilha", type=Path, help=".xlsx (ou .csv antigo)")
    a = ap.parse_args()
    if a.cmd == "gerar":
        gerar(a.jsonl, a.repeticao)
    elif a.cmd == "converter":
        converter(a.csv)
    else:
        revelar(a.planilha)


if __name__ == "__main__":
    main()
